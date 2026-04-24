from docx import Document

# Create a reference DOCX with preferred table style
doc = Document()
doc.add_paragraph('Reference document for pandoc to apply table styles')
# create a sample table to ensure styles are embedded
table = doc.add_table(rows=2, cols=2)
for i, row in enumerate(table.rows):
    for j, cell in enumerate(row.cells):
        cell.text = f'Sample {i},{j}'
# Try to set a built-in style that tends to be present in Word
try:
    table.style = 'Light Shading'
except Exception:
    try:
        table.style = 'Table Grid'
    except Exception:
        pass

doc.save('docs/reference.docx')
print('Wrote docs/reference.docx')
