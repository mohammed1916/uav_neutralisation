import os
import zipfile
from docx import Document

DOC = os.path.normpath(os.path.join(os.path.dirname(__file__), '..', 'docs', 'pneumatic_launcher_analysis_complete.docx'))

print('Document:', DOC)
if not os.path.exists(DOC):
    print('File not found.')
else:
    # basic docx introspection (paragraphs, tables, images)
    try:
        doc = Document(DOC)
        print('Paragraphs:', len(doc.paragraphs))
        heads = [p.text for p in doc.paragraphs if p.style.name.startswith('Heading') and p.text.strip()]
        print('Headings found:')
        for h in heads:
            print('-', h)
        print('Tables found:', len(doc.tables))
        for i, t in enumerate(doc.tables):
            print(f'\nTable {i+1}: rows={len(t.rows)} cols={len(t.columns)}')
            for r in t.rows[:5]:
                cells = [c.text.strip() for c in r.cells]
                print(' | '.join(cells))
        rels = doc.part._rels
        imgs = []
        for r in rels:
            rel = rels[r]
            if 'image' in getattr(rel, 'target_ref', ''):
                imgs.append(rel.target_ref)
        print('Embedded image parts:')
        for im in imgs:
            print('-', im)
    except Exception:
        pass

    # Detect OMML math by reading word/document.xml directly
    omml_present = False
    try:
        with zipfile.ZipFile(DOC, 'r') as z:
            if 'word/document.xml' in z.namelist():
                data = z.read('word/document.xml')
                # look for OMML markers
                if b'<m:oMath' in data or b'<m:oMathPara' in data:
                    omml_present = True
    except Exception:
        omml_present = False

    print('OMML math present:', omml_present)
